"""Módulo para la exportación de datos de facturación a formato Excel (.xlsx).

Este módulo utiliza la librería openpyxl para generar archivos de Excel siguiendo
un formato estándar corporativo (Quipux SAS). Incluye configuraciones de estilos,
encabezados combinados, bordes y formatos de fuente específicos.
"""

import os
import yaml
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Fill, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

class ExcelExporter:
    """Clase encargada de la generación de archivos Excel para reportes de control.
    
    Carga la configuración de columnas y estilos desde un archivo de metadatos
    y aplica el formato requerido a las hojas de cálculo.
    """

    def __init__(self):
        """Inicializa el exportador cargando la configuración desde excel_config.yml."""
        config_path = os.path.join("metadata", "excel_config.yml")
        with open(config_path, "r", encoding="utf-8") as f:
            self.config = yaml.safe_load(f)
        
        self.styles = self.config["styles"]
        self.columns = self.config["columns"]

    def generate_excel(
        self,
        data: list[dict],
        month_name: str,
        year: str | int,
    ) -> Workbook:
        """Genera un objeto Workbook con los datos y formatos especificados.

        Args:
            data: Lista de diccionarios con los registros de la base de datos.
            month_name: Nombre del mes o rango de fechas para el encabezado.
            year: Año del reporte para el encabezado.

        Returns:
            Workbook listo para ser guardado o transmitido.

        """
        wb = Workbook()
        ws = wb.active
        ws.title = self.config.get("sheet_name", "Facturas")

        # Styles
        header_fill = PatternFill(start_color=self.styles["header_bg_color"], 
                                  end_color=self.styles["header_bg_color"], 
                                  fill_type="solid")
        header_font = Font(name=self.styles["font_name"], 
                           size=self.styles["font_size"], 
                           bold=True, 
                           color=self.styles["header_font_color"])
        
        # Border styles
        medium_side = Side(style='medium', color='000000')
        thin_side = Side(style='thin', color='000000')
        
        thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        # 1. Quipux SAS Header
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(self.columns))
        cell_1 = ws.cell(row=1, column=1, value=self.config["company_name"].upper())
        cell_1.font = Font(name=self.styles["font_name"], size=12, bold=True)
        cell_1.alignment = Alignment(horizontal="center")
        # Apply borders to Row 1
        for i in range(1, len(self.columns) + 1):
            cell = ws.cell(row=1, column=i)
            # Special case for column 1, 2 and last
            r_border = medium_side if i in (1, 2, len(self.columns)) else thin_side
            l_border = medium_side if i == 1 else (medium_side if i in (2, 3) else thin_side)
            # Simplifiying: User says: "borde derecho de la columna EVENTO DIAN ... y asi para la columna 1 y 2"
            # And "cuadricula bordes de 1.0pt"
            right_s = medium_side if i in (1, 2, len(self.columns)) else thin_side
            left_s = medium_side if i == 1 else (medium_side if i in (2, 3) else thin_side) # If i=2, its left is i=1's right
            cell.border = Border(left=thin_side, right=right_s, top=thin_side, bottom=thin_side)

        # 2. Title Header
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(self.columns))
        title = self.config["report_title"].format(MONTH=month_name.upper(), YEAR=year).upper()
        cell_2 = ws.cell(row=2, column=1, value=title)
        cell_2.font = Font(name=self.styles["font_name"], size=11, bold=True)
        cell_2.alignment = Alignment(horizontal="center")
        for i in range(1, len(self.columns) + 1):
            cell = ws.cell(row=2, column=i)
            right_s = medium_side if i in (1, 2, len(self.columns)) else thin_side
            cell.border = Border(left=thin_side, right=right_s, top=thin_side, bottom=thin_side)

        # 3. Table Headers
        for i, col in enumerate(self.columns, 1):
            cell = ws.cell(row=3, column=i, value=col["label"])
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            right_s = medium_side if i in (1, 2, len(self.columns)) else thin_side
            cell.border = Border(left=thin_side, right=right_s, top=thin_side, bottom=thin_side)

        # 4. Data Rows
        num_rows = len(data)
        for r_idx, row_data in enumerate(data, 4):
            is_last_row = (r_idx == num_rows + 3)
            for c_idx, col in enumerate(self.columns, 1):
                field = col["db_field"]
                value = row_data.get(field)

                # Transform boolean to 'X'
                if field in ["acuso_recibido", "recibido_bien_servicio", "aceptacion_empresa", "recibido"]:
                    value = "X" if value is True else ""
                
                # Format dates
                if isinstance(value, datetime):
                    value = value.strftime("%d/%m/%Y")
                elif value is None:
                    value = ""
                
                # Convert to uppercase string
                value = str(value).upper()
                
                cell = ws.cell(row=r_idx, column=c_idx, value=value)
                cell.font = Font(name=self.styles["font_name"], size=self.styles["font_size"])
                
                # Borders
                right_s = medium_side if c_idx in (1, 2, len(self.columns)) else thin_side
                bottom_s = medium_side if is_last_row else thin_side
                cell.border = Border(left=thin_side, right=right_s, top=thin_side, bottom=bottom_s)
                
                # Alignment
                h_align = "left"
                # c_idx 1: Fecha em, 3: Fecha ent, 5: NIT, 7: No. Factura
                if c_idx in (1, 3, 5, 7):
                    h_align = "right"
                
                cell.alignment = Alignment(horizontal=h_align, vertical="center")

        return wb
